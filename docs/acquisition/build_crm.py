from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

F = "Arial"
hdr_fill = PatternFill("solid", start_color="1F4E78")
hdr_font = Font(name=F, bold=True, color="FFFFFF", size=10)
score_fill = PatternFill("solid", start_color="FFF2CC")
thin = Side(style="thin", color="D9D9D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()

# ---- Sheet 1: Leady ----
ws = wb.active
ws.title = "Leady"
cols = [
 ("Firma",22),("Web",26),("Obor",16),("Město",14),("Zdroj",16),
 ("Problém 1",24),("Problém 2",24),("Problém 3",24),
 ("Viditelný problém (0-3)",12),("Ruční administrativa (0-3)",12),
 ("Schopnost platit (0-2)",11),("Rozhodovatel/kanál (0-1)",11),("Referral cesta (0-1)",11),
 ("SKÓRE CELKEM",13),
 ("Doporučená služba",24),("Kanál",14),("Datum oslovení",13),("Reakce",14),
 ("Další krok",24),("Stav",16),("Poznámka",30),
]
for i,(name,w) in enumerate(cols,1):
    c = ws.cell(1,i,name); c.font=hdr_font; c.fill=hdr_fill; c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center"); c.border=border
    ws.column_dimensions[ws.cell(1,i).column_letter].width = w
ws.row_dimensions[1].height = 42
ws.freeze_panes = "A2"

ROWS = 60
for r in range(2, ROWS+2):
    # total score formula = sum of the 5 scoring cols (I:M)
    tc = ws.cell(r,14, f"=SUM(I{r}:M{r})")
    tc.font=Font(name=F,bold=True); tc.alignment=Alignment(horizontal="center")
    for i in range(9,14):
        ws.cell(r,i).fill=score_fill; ws.cell(r,i).alignment=Alignment(horizontal="center")
    for i in range(1,22):
        ws.cell(r,i).border=border
        if ws.cell(r,i).font.name is None: ws.cell(r,i).font=Font(name=F,size=10)

# dropdowns
dv_stav = DataValidation(type="list", formula1='"nalezeno,audit připraven,osloveno,odpověď,call,nabídka,vyhráno,ztraceno,nekontaktovat"', allow_blank=True)
dv_kanal = DataValidation(type="list", formula1='"reaktivace,partner/referral,mini-audit,formulář,LinkedIn,radar/portál"', allow_blank=True)
ws.add_data_validation(dv_stav); ws.add_data_validation(dv_kanal)
dv_stav.add(f"T2:T{ROWS+1}"); dv_kanal.add(f"P2:P{ROWS+1}")

# highlight SKÓRE CELKEM >=7 green, <=4 red
ws.conditional_formatting.add(f"N2:N{ROWS+1}", CellIsRule(operator="greaterThanOrEqual", formula=["7"], fill=PatternFill("solid",start_color="C6EFCE"), font=Font(name=F,bold=True,color="006100")))
ws.conditional_formatting.add(f"N2:N{ROWS+1}", CellIsRule(operator="lessThanOrEqual", formula=["4"], fill=PatternFill("solid",start_color="FFC7CE"), font=Font(name=F,color="9C0006")))

# 2 example rows
ex = [
 ["Truhlářství Novák","trunovak.cz","truhlářství","Kladno","Mapy.cz","mobil bez kontaktu","8s načítání","žádné CTA",3,1,1,1,0,None,"oprava mobilu + CTA","mini-audit","","","poslat 3 body","audit připraven","skóre 6 – až po 7+"],
 ["VO Stavebniny s.r.o.","stavmat-vo.cz","velkoobchod","Brno","referral (účetní)","objednávky e-mailem ručně do Pohody","—","—",1,3,2,1,1,None,"prototyp automatizace objednávek","partner/referral","","","call 15 min","osloveno","HORKÝ – automatizace"],
]
for ri,row in enumerate(ex,2):
    for ci,val in enumerate(row,1):
        if ci==14: continue
        ws.cell(ri,ci,val)

# ---- Sheet 2: Návod ----
ns = wb.create_sheet("Návod & skórování")
ns.column_dimensions["A"].width=42; ns.column_dimensions["B"].width=10
def h(r,t):
    c=ns.cell(r,1,t); c.font=Font(name=F,bold=True,size=12,color="1F4E78")
def line(r,a,b=None):
    ns.cell(r,1,a).font=Font(name=F,size=10)
    if b is not None: ns.cell(r,2,b).font=Font(name=F,size=10,bold=True)
h(1,"Skórování leadu (oslovuj přednostně 7+)")
rubric=[("Viditelný problém na webu/formuláři","0–3"),("Pravděpodobná ruční administrativa","0–3"),
 ("Reálná schopnost platit","0–2"),("Rozhodovatel / dobrý kontaktní kanál","0–1"),("Referral cesta přes partnera","0–1"),("MAX","10")]
for i,(a,b) in enumerate(rubric,2): line(i,a,b)
h(9,"Stavy (sloupec Stav)")
for i,s in enumerate(["nalezeno → audit připraven → osloveno → odpověď → call → nabídka → vyhráno","ztraceno / nekontaktovat"],10): line(i,s)
h(13,"Cíl po 14 dnech")
for i,s in enumerate(["15 reaktivačních zpráv","5–10 partnerských zpráv","~20 mini-auditů / oslovení","→ realisticky 3–8 odpovědí, 1–3 cally, 1 placená kontrola"],14): line(i,s)
h(19,"Pravidlo")
for i,s in enumerate(["0 odpovědí z 20 konkrétních oslovení = problém je cílení/zpráva/nabídka, ne 'málo reklamy'.","Scraper/automatizaci stav až po 30–50 ručních osloveních (fáze 2)."],20): line(i,s)

wb.save("CRM-leady-radeq.xlsx")
print("saved")
